from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q, Sum
from django.shortcuts import (
    get_object_or_404,
    redirect,
    render,
)
from django.contrib.auth import get_user_model, update_session_auth_hash
from django.core.exceptions import PermissionDenied
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import (
    ProductForm,
    DriverForm,
    TruckForm,
    DistributorForm,
    TerritoryForm,
    OrderForm,
    OrderItemFormSet,
    DispatchForm,
    DispatchItemFormSet,
    DeliveryNoteForm,
    OMSUserCreationForm,
    OMSUserEditForm,
    OMSUserPasswordForm,
)

from .models import (
    Territory,
    Distributor,
    Product,
    Driver,
    Truck,
    Order,
    OrderItem,
    Dispatch,
    DispatchItem,
    PurchaseRequest,
    DeliveryNote,
)

from io import BytesIO

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def _generate_order_serial(order_date):
    """
    Generate an order serial number using:

        YYMMDD + daily sequence

    Example:
        26081701
        26081702

    The user can still edit the generated serial number.
    """

    prefix = order_date.strftime("%y%m%d")

    existing_serials = (
        Order.objects
        .filter(serial_number__startswith=prefix)
        .values_list("serial_number", flat=True)
    )

    sequence_numbers = []

    for serial in existing_serials:

        # Standard serial:
        # YYMMDDXX = 8 characters
        if (
            len(serial) == 8
            and serial[:6] == prefix
            and serial[6:].isdigit()
        ):
            sequence_numbers.append(
                int(serial[6:])
            )

    next_number = max(
        sequence_numbers,
        default=0
    ) + 1

    return f"{prefix}{next_number:02d}"


def _generate_dispatch_number(order):
    """
    Suggest a dispatch number.

    First dispatch:
        same as order serial

    Later dispatches:
        SERIAL-2
        SERIAL-3
        etc.

    The clerk can still edit it.
    """

    dispatch_count = order.dispatches.count()

    if dispatch_count == 0:
        return order.serial_number

    return (
        f"{order.serial_number}-"
        f"{dispatch_count + 1}"
    )


def _validate_dispatch_quantities(
    formset,
    order,
    dispatch=None,
):
    """
    Additional validation for dispatch quantities.

    Ensures:

    1. Every DispatchItem belongs to the order.
    2. Total dispatched quantity does not exceed
       the original ordered quantity.

    This is especially important while creating a
    new Dispatch because the Dispatch object has not
    yet been saved to the database.
    """

    valid = True

    for form in formset.forms:

        if not hasattr(form, "cleaned_data"):
            continue

        if form.cleaned_data.get("DELETE"):
            continue

        order_item = form.cleaned_data.get(
            "order_item"
        )

        quantity = form.cleaned_data.get(
            "quantity"
        )

        if not order_item or not quantity:
            continue

        # Make sure the selected item belongs
        # to this order.
        if order_item.order_id != order.id:

            form.add_error(
                "order_item",
                "This product does not belong "
                "to this order."
            )

            valid = False
            continue

        previous_dispatches = (
            DispatchItem.objects
            .filter(order_item=order_item)
        )

        # When editing a dispatch, don't count
        # its current quantities twice.
        if dispatch and dispatch.pk:

            previous_dispatches = (
                previous_dispatches
                .exclude(dispatch=dispatch)
            )

        already_dispatched = (
            previous_dispatches
            .aggregate(
                total=Sum("quantity")
            )["total"]
            or 0
        )

        if (
            already_dispatched + quantity
            > order_item.quantity
        ):

            remaining = (
                order_item.quantity
                - already_dispatched
            )

            form.add_error(
                "quantity",
                (
                    "Quantity exceeds the amount "
                    "remaining on the order. "
                    f"Remaining: {remaining} cases."
                )
            )

            valid = False

    return valid

def _get_dispatch_order_item_statuses(
    order,
    dispatch=None,
):
    """
    Build quantity information for every product
    belonging to an Order.

    For a new Dispatch:

        previously_dispatched =
            quantity in all previous Dispatches

        available =
            ordered - previously_dispatched

    When editing a Dispatch, quantities belonging to
    the Dispatch currently being edited are excluded
    from "previously dispatched".

    This means the current Dispatch can still be
    corrected without counting itself twice.
    """

    order_items = list(
        order.items
        .select_related("product")
        .order_by(
            "product__name",
            "product__size",
        )
    )

    dispatched_items = (
        DispatchItem.objects
        .filter(
            order_item__order=order
        )
    )

    if dispatch and dispatch.pk:

        dispatched_items = (
            dispatched_items
            .exclude(
                dispatch=dispatch
            )
        )

    dispatched_totals = {
        row["order_item_id"]:
            row["total"] or 0

        for row in (
            dispatched_items
            .values("order_item_id")
            .annotate(
                total=Sum("quantity")
            )
        )
    }

    statuses = []

    for order_item in order_items:

        previously_dispatched = (
            dispatched_totals.get(
                order_item.pk,
                0,
            )
        )

        available = (
            order_item.quantity
            - previously_dispatched
        )

        if available < 0:
            available = 0

        statuses.append(
            {
                "order_item": order_item,
                "ordered": order_item.quantity,
                "previously_dispatched":
                    previously_dispatched,
                "available": available,
            }
        )

    return statuses


def _attach_dispatch_form_metadata(
    formset,
    statuses,
):
    """
    Attach quantity information to each DispatchItem
    form so the template can display:

        Ordered
        Previously Dispatched
        Available

    This keeps quantity calculations in Python rather
    than duplicating business logic in JavaScript.
    """

    status_map = {
        status["order_item"].pk:
            status

        for status in statuses
    }

    for item_form in formset.forms:

        order_item_id = None

        # ---------------------------------------------
        # Bound POST form
        # ---------------------------------------------

        if item_form.is_bound:

            raw_order_item_id = (
                item_form.data.get(
                    item_form.add_prefix(
                        "order_item"
                    )
                )
            )

            if raw_order_item_id:

                try:
                    order_item_id = int(
                        raw_order_item_id
                    )

                except (
                    TypeError,
                    ValueError,
                ):
                    order_item_id = None

        # ---------------------------------------------
        # Existing DispatchItem
        # ---------------------------------------------

        if (
            order_item_id is None
            and item_form.instance
            and item_form.instance.pk
        ):

            order_item_id = (
                item_form.instance
                .order_item_id
            )

        # ---------------------------------------------
        # Initial form
        # ---------------------------------------------

        if order_item_id is None:

            initial_order_item = (
                item_form.initial.get(
                    "order_item"
                )
            )

            if hasattr(
                initial_order_item,
                "pk",
            ):

                order_item_id = (
                    initial_order_item.pk
                )

            elif initial_order_item:

                try:
                    order_item_id = int(
                        initial_order_item
                    )

                except (
                    TypeError,
                    ValueError,
                ):
                    order_item_id = None

        metadata = status_map.get(
            order_item_id
        )

        item_form.dispatch_meta = (
            metadata
        )

        if metadata:

            item_form.fields[
                "quantity"
            ].widget.attrs["max"] = (
                metadata["available"]
            )


def _parse_date(value):
    """
    Safely convert YYYY-MM-DD from query strings
    into a Python date.
    """

    if not value:
        return None

    try:
        return datetime.strptime(
            value,
            "%Y-%m-%d"
        ).date()

    except ValueError:
        return None
    

def _generate_driver_code():
    """
    Generate the next Driver ID.

    Format:
        DRV0001
        DRV0002
        DRV0003

    Existing codes that do not follow this format
    are ignored.
    """

    prefix = "DRV"

    existing_codes = (
        Driver.objects
        .filter(
            driver_code__startswith=prefix
        )
        .values_list(
            "driver_code",
            flat=True,
        )
    )

    numbers = []

    for code in existing_codes:

        suffix = code[
            len(prefix):
        ]

        if suffix.isdigit():
            numbers.append(
                int(suffix)
            )

    next_number = (
        max(
            numbers,
            default=0,
        )
        + 1
    )

    # Extra protection in case a matching
    # code already exists.
    while True:

        candidate = (
            f"{prefix}"
            f"{next_number:04d}"
        )

        if not Driver.objects.filter(
            driver_code=candidate
        ).exists():
            return candidate

        next_number += 1

def _generate_customer_code():
    """
    Generate the next Customer ID.

    Format:
        CUST000001
        CUST000002
        CUST000003

    Existing customer codes that do not follow
    this format are ignored.
    """

    prefix = "CUST"

    existing_codes = (
        Distributor.objects
        .filter(
            customer_code__startswith=prefix
        )
        .values_list(
            "customer_code",
            flat=True,
        )
    )

    numbers = []

    for code in existing_codes:

        suffix = code[
            len(prefix):
        ]

        if suffix.isdigit():
            numbers.append(
                int(suffix)
            )

    next_number = (
        max(
            numbers,
            default=0,
        )
        + 1
    )

    while True:

        candidate = (
            f"{prefix}"
            f"{next_number:06d}"
        )

        if not Distributor.objects.filter(
            customer_code=candidate
        ).exists():
            return candidate

        next_number += 1
# =========================================================
# USER MANAGEMENT
# =========================================================


User = get_user_model()


def _require_superuser(request):
    """
    Restrict OMS user-management functions to
    Django superusers / system administrators.
    """

    if not request.user.is_superuser:
        raise PermissionDenied



# =========================================================
# USER LIST
# =========================================================

@login_required
def user_list(request):

    _require_superuser(request)


    users = (
        User.objects
        .all()
        .order_by(
            "username"
        )
    )


    search = (
        request.GET
        .get(
            "q",
            ""
        )
        .strip()
    )


    status = (
        request.GET
        .get(
            "status",
            ""
        )
        .strip()
    )


    # =====================================================
    # SEARCH
    # =====================================================

    if search:

        users = users.filter(
            Q(
                username__icontains=
                    search
            )
            |
            Q(
                first_name__icontains=
                    search
            )
            |
            Q(
                last_name__icontains=
                    search
            )
            |
            Q(
                email__icontains=
                    search
            )
        )


    # =====================================================
    # STATUS FILTER
    # =====================================================

    if status == "active":

        users = users.filter(
            is_active=True
        )


    elif status == "inactive":

        users = users.filter(
            is_active=False
        )


    context = {
        "system_users":
            users,

        "search":
            search,

        "selected_status":
            status,
    }


    return render(
        request,
        "orders/user_list.html",
        context,
    )



# =========================================================
# CREATE USER
# =========================================================

@login_required
def user_create(request):

    _require_superuser(request)


    if request.method == "POST":

        form = OMSUserCreationForm(
            request.POST
        )


        if form.is_valid():

            new_user = (
                form.save()
            )


            messages.success(
                request,
                (
                    f"User "
                    f"{new_user.username} "
                    f"was created successfully."
                )
            )


            return redirect(
                "orders:user_list"
            )


    else:

        form = OMSUserCreationForm(
            initial={
                "is_active": True,
            }
        )


    context = {
        "form":
            form,

        "page_title":
            "Add User",

        "editing":
            False,
    }


    return render(
        request,
        "orders/user_form.html",
        context,
    )



# =========================================================
# EDIT USER
# =========================================================

@login_required
def user_edit(
    request,
    user_id,
):

    _require_superuser(request)


    managed_user = get_object_or_404(
        User,
        pk=user_id,
    )


    if request.method == "POST":

        form = OMSUserEditForm(
            request.POST,
            instance=managed_user,
        )


        if form.is_valid():

            updated_user = (
                form.save(
                    commit=False
                )
            )


            # =================================================
            # DO NOT ALLOW CURRENT USER TO DEACTIVATE THEMSELF
            # =================================================

            if (
                updated_user.pk
                ==
                request.user.pk
                and
                not updated_user.is_active
            ):

                messages.error(
                    request,
                    (
                        "You cannot deactivate "
                        "your own account."
                    )
                )


            # =================================================
            # PROTECT LAST ACTIVE ADMINISTRATOR
            # =================================================

            elif (
                updated_user.is_superuser
                and
                not updated_user.is_active
            ):

                other_active_admin_exists = (
                    User.objects
                    .filter(
                        is_superuser=True,
                        is_active=True,
                    )
                    .exclude(
                        pk=updated_user.pk
                    )
                    .exists()
                )


                if not other_active_admin_exists:

                    messages.error(
                        request,
                        (
                            "The last active administrator "
                            "cannot be deactivated."
                        )
                    )


                else:

                    updated_user.save()


                    messages.success(
                        request,
                        (
                            f"User "
                            f"{updated_user.username} "
                            f"was updated successfully."
                        )
                    )


                    return redirect(
                        "orders:user_list"
                    )


            # =================================================
            # NORMAL UPDATE
            # =================================================

            else:

                updated_user.save()


                messages.success(
                    request,
                    (
                        f"User "
                        f"{updated_user.username} "
                        f"was updated successfully."
                    )
                )


                return redirect(
                    "orders:user_list"
                )


    else:

        form = OMSUserEditForm(
            instance=managed_user
        )


    context = {
        "form":
            form,

        "managed_user":
            managed_user,

        "page_title":
            "Edit User",

        "editing":
            True,
    }


    return render(
        request,
        "orders/user_form.html",
        context,
    )



# =========================================================
# ACTIVATE / DEACTIVATE USER
# =========================================================

@login_required
def user_toggle_status(
    request,
    user_id,
):

    _require_superuser(request)


    if request.method != "POST":

        return redirect(
            "orders:user_list"
        )


    managed_user = get_object_or_404(
        User,
        pk=user_id,
    )


    # =====================================================
    # CURRENT USER PROTECTION
    # =====================================================

    if (
        managed_user.pk
        ==
        request.user.pk
    ):

        messages.error(
            request,
            (
                "You cannot deactivate "
                "your own account."
            )
        )


        return redirect(
            "orders:user_list"
        )


    # =====================================================
    # DEACTIVATE USER
    # =====================================================

    if managed_user.is_active:


        # -------------------------------------------------
        # Protect the final active administrator
        # -------------------------------------------------

        if managed_user.is_superuser:

            other_active_admin_exists = (
                User.objects
                .filter(
                    is_superuser=True,
                    is_active=True,
                )
                .exclude(
                    pk=managed_user.pk
                )
                .exists()
            )


            if not other_active_admin_exists:

                messages.error(
                    request,
                    (
                        "The last active administrator "
                        "cannot be deactivated."
                    )
                )


                return redirect(
                    "orders:user_list"
                )


        managed_user.is_active = False


        managed_user.save(
            update_fields=[
                "is_active"
            ]
        )


        messages.success(
            request,
            (
                f"User "
                f"{managed_user.username} "
                f"was deactivated."
            )
        )


    # =====================================================
    # ACTIVATE USER
    # =====================================================

    else:

        managed_user.is_active = True


        managed_user.save(
            update_fields=[
                "is_active"
            ]
        )


        messages.success(
            request,
            (
                f"User "
                f"{managed_user.username} "
                f"was activated."
            )
        )


    return redirect(
        "orders:user_list"
    )



# =========================================================
# RESET USER PASSWORD
# =========================================================

@login_required
def user_reset_password(
    request,
    user_id,
):

    _require_superuser(request)


    managed_user = get_object_or_404(
        User,
        pk=user_id,
    )


    if request.method == "POST":

        form = OMSUserPasswordForm(
            managed_user,
            request.POST,
        )


        if form.is_valid():

            updated_user = (
                form.save()
            )


            # -------------------------------------------------
            # If an administrator changes their own password,
            # keep the current login session valid.
            # -------------------------------------------------

            if (
                updated_user.pk
                ==
                request.user.pk
            ):

                update_session_auth_hash(
                    request,
                    updated_user,
                )


            messages.success(
                request,
                (
                    f"Password for "
                    f"{updated_user.username} "
                    f"was changed successfully."
                )
            )


            return redirect(
                "orders:user_list"
            )


    else:

        form = OMSUserPasswordForm(
            managed_user
        )


    context = {
        "form":
            form,

        "managed_user":
            managed_user,
    }


    return render(
        request,
        "orders/user_password_form.html",
        context,
    )



# =========================================================
# DASHBOARD
# =========================================================

@login_required
def dashboard(request):

    recent_orders = (
        Order.objects
        .select_related(
            "distributor",
            "distributor__territory"
        )
        .order_by("-created_at")[:10]
    )

    context = {
        "active_territories":
            Territory.objects.filter(
                is_active=True
            ).count(),

        "active_distributors":
            Distributor.objects.filter(
                is_active=True
            ).count(),

        "active_products":
            Product.objects.filter(
                is_active=True
            ).count(),

        "active_drivers":
            Driver.objects.filter(
                is_active=True
            ).count(),

        "active_trucks":
            Truck.objects.filter(
                is_active=True
            ).count(),

        "total_orders":
            Order.objects.count(),

        "today_orders":
            Order.objects.filter(
                order_date=timezone.localdate()
            ).count(),

        "recent_orders":
            recent_orders,
    }

    return render(
        request,
        "orders/dashboard.html",
        context,
    )


# =========================================================
# HELP
# =========================================================


@login_required
def help_page(request):

    return render(
        request,
        "orders/help.html",
    )


# =========================================================
# MANAGEMENT
# =========================================================

@login_required
def management_home(request):
    context = {
        "product_count": Product.objects.count(),
        "active_product_count": Product.objects.filter(is_active=True).count(),
        "driver_count": Driver.objects.count(),
        "active_driver_count": Driver.objects.filter(is_active=True).count(),
        "truck_count": Truck.objects.count(),
        "active_truck_count": Truck.objects.filter(is_active=True).count(),
        "distributor_count": Distributor.objects.count(),
        "active_distributor_count": Distributor.objects.filter(is_active=True).count(),
        "territory_count": Territory.objects.count(),
        "active_territory_count": Territory.objects.filter(is_active=True).count(),
    }

    return render(
        request,
        "orders/management_home.html",
        context,
    )


# ---------------------------------------------------------
# PRODUCTS
# ---------------------------------------------------------

@login_required
def product_list(request):
    products = Product.objects.all().order_by("name", "size")

    search = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    package_type = request.GET.get("package_type", "").strip()

    if search:
        products = products.filter(
            Q(sku__icontains=search)
            | Q(name__icontains=search)
            | Q(size__icontains=search)
        )

    if status == "active":
        products = products.filter(is_active=True)
    elif status == "inactive":
        products = products.filter(is_active=False)

    if package_type:
        products = products.filter(package_type=package_type)

    context = {
        "products": products,
        "search": search,
        "selected_status": status,
        "selected_package_type": package_type,
        "package_choices": Product.PACKAGE_TYPES,
    }

    return render(request, "orders/product_list.html", context)


@login_required
def product_create(request):
    if request.method == "POST":
        form = ProductForm(request.POST)

        if form.is_valid():
            product = form.save()
            messages.success(
                request,
                f"Product {product.sku} was added successfully."
            )
            return redirect("orders:product_list")
    else:
        form = ProductForm(initial={"is_active": True})

    return render(
        request,
        "orders/product_form.html",
        {
            "form": form,
            "page_title": "Add Product",
            "editing": False,
        },
    )


@login_required
def product_edit(request, product_id):
    product = get_object_or_404(Product, pk=product_id)

    if request.method == "POST":
        form = ProductForm(request.POST, instance=product)

        if form.is_valid():
            product = form.save()
            messages.success(
                request,
                f"Product {product.sku} was updated successfully."
            )
            return redirect("orders:product_list")
    else:
        form = ProductForm(instance=product)

    return render(
        request,
        "orders/product_form.html",
        {
            "form": form,
            "product": product,
            "page_title": "Edit Product",
            "editing": True,
        },
    )


@login_required
@require_POST
def product_toggle_status(request, product_id):
    product = get_object_or_404(Product, pk=product_id)
    product.is_active = not product.is_active
    product.save(update_fields=["is_active"])

    state = "activated" if product.is_active else "deactivated"
    messages.success(request, f"Product {product.sku} was {state}.")
    return redirect("orders:product_list")


# ---------------------------------------------------------
# DRIVERS
# ---------------------------------------------------------

@login_required
def driver_list(request):

    drivers = (
        Driver.objects
        .all()
        .order_by(
            "driver_code"
        )
    )

    search = request.GET.get(
        "q",
        ""
    ).strip()

    status = request.GET.get(
        "status",
        ""
    ).strip()


    # =====================================================
    # SEARCH
    # =====================================================

    if search:

        drivers = drivers.filter(
            Q(
                driver_code__icontains=
                search
            )
            |
            Q(
                full_name__icontains=
                search
            )
        )


    # =====================================================
    # STATUS
    # =====================================================

    if status == "active":

        drivers = drivers.filter(
            is_active=True
        )

    elif status == "inactive":

        drivers = drivers.filter(
            is_active=False
        )


    context = {
        "drivers":
            drivers,

        "search":
            search,

        "selected_status":
            status,
    }


    return render(
        request,
        "orders/driver_list.html",
        context,
    )


@login_required
def driver_create(request):

    if request.method == "POST":

        data = request.POST.copy()

        # If the field is submitted empty,
        # generate an ID on the server.
        if not data.get(
            "driver_code"
        ):

            data["driver_code"] = (
                _generate_driver_code()
            )

        form = DriverForm(
            data
        )

        if form.is_valid():

            driver = form.save()

            messages.success(
                request,
                (
                    f"Driver "
                    f"{driver.full_name} "
                    f"({driver.driver_code}) "
                    f"was created successfully."
                )
            )

            return redirect(
                "orders:driver_list"
            )

    else:

        form = DriverForm(
            initial={
                "driver_code":
                    _generate_driver_code(),

                "is_active":
                    True,
            }
        )


    context = {
        "form":
            form,

        "page_title":
            "Add Driver",

        "editing":
            False,
    }


    return render(
        request,
        "orders/driver_form.html",
        context,
    )

@login_required
def driver_edit(request, driver_id):
    driver = get_object_or_404(Driver, pk=driver_id)

    if request.method == "POST":
        form = DriverForm(request.POST, instance=driver)
        if form.is_valid():
            driver = form.save()
            messages.success(
                request,
                f"Driver {driver.full_name} was updated successfully."
            )
            return redirect("orders:driver_list")
    else:
        form = DriverForm(instance=driver)

    return render(
        request,
        "orders/driver_form.html",
        {
            "form": form,
            "driver": driver,
            "page_title": "Edit Driver",
            "editing": True,
        },
    )


@login_required
@require_POST
def driver_toggle_status(request, driver_id):
    driver = get_object_or_404(Driver, pk=driver_id)
    driver.is_active = not driver.is_active
    driver.save(update_fields=["is_active"])

    state = "activated" if driver.is_active else "deactivated"
    messages.success(request, f"Driver {driver.full_name} was {state}.")
    return redirect("orders:driver_list")


# ---------------------------------------------------------
# TRUCKS
# ---------------------------------------------------------

@login_required
def truck_list(request):

    trucks = (
        Truck.objects
        .all()
        .order_by(
            "plate_number"
        )
    )

    search = request.GET.get(
        "q",
        ""
    ).strip()

    status = request.GET.get(
        "status",
        ""
    ).strip()


    # =====================================================
    # SEARCH
    # =====================================================

    if search:

        trucks = trucks.filter(
            Q(
                plate_number__icontains=
                search
            )
            |
            Q(
                model__icontains=
                search
            )
        )


    # =====================================================
    # STATUS
    # =====================================================

    if status == "active":

        trucks = trucks.filter(
            is_active=True
        )

    elif status == "inactive":

        trucks = trucks.filter(
            is_active=False
        )


    context = {
        "trucks":
            trucks,

        "search":
            search,

        "selected_status":
            status,
    }


    return render(
        request,
        "orders/truck_list.html",
        context,
    )

@login_required
def truck_create(request):
    if request.method == "POST":
        form = TruckForm(request.POST)
        if form.is_valid():
            truck = form.save()
            messages.success(
                request,
                f"Truck {truck.plate_number} was added successfully."
            )
            return redirect("orders:truck_list")
    else:
        form = TruckForm(initial={"is_active": True})

    return render(
        request,
        "orders/truck_form.html",
        {"form": form, "page_title": "Add Truck", "editing": False},
    )


@login_required
def truck_edit(request, truck_id):
    truck = get_object_or_404(Truck, pk=truck_id)

    if request.method == "POST":
        form = TruckForm(request.POST, instance=truck)
        if form.is_valid():
            truck = form.save()
            messages.success(
                request,
                f"Truck {truck.plate_number} was updated successfully."
            )
            return redirect("orders:truck_list")
    else:
        form = TruckForm(instance=truck)

    return render(
        request,
        "orders/truck_form.html",
        {
            "form": form,
            "truck": truck,
            "page_title": "Edit Truck",
            "editing": True,
        },
    )


@login_required
@require_POST
def truck_toggle_status(request, truck_id):
    truck = get_object_or_404(Truck, pk=truck_id)
    truck.is_active = not truck.is_active
    truck.save(update_fields=["is_active"])

    state = "activated" if truck.is_active else "deactivated"
    messages.success(request, f"Truck {truck.plate_number} was {state}.")
    return redirect("orders:truck_list")


# ---------------------------------------------------------
# CUSTOMERS / DISTRIBUTORS
# ---------------------------------------------------------

@login_required
def distributor_list(request):

    distributors = (
        Distributor.objects
        .select_related(
            "territory"
        )
        .all()
        .order_by(
            "full_name"
        )
    )

    search = request.GET.get(
        "q",
        ""
    ).strip()

    status = request.GET.get(
        "status",
        ""
    ).strip()

    customer_type = request.GET.get(
        "customer_type",
        ""
    ).strip()

    payment_status = request.GET.get(
        "payment_status",
        ""
    ).strip()

    territory_id = request.GET.get(
        "territory",
        ""
    ).strip()


    # =====================================================
    # SEARCH
    # =====================================================

    if search:

        distributors = (
            distributors.filter(
                Q(
                    customer_code__icontains=
                    search
                )
                |
                Q(
                    full_name__icontains=
                    search
                )
                |
                Q(
                    territory__name__icontains=
                    search
                )
            )
        )


    # =====================================================
    # STATUS
    # =====================================================

    if status == "active":

        distributors = (
            distributors.filter(
                is_active=True
            )
        )

    elif status == "inactive":

        distributors = (
            distributors.filter(
                is_active=False
            )
        )


    # =====================================================
    # CUSTOMER TYPE
    # =====================================================

    valid_customer_types = {
        value
        for value, label
        in Distributor.CUSTOMER_TYPES
    }

    if customer_type in valid_customer_types:

        distributors = (
            distributors.filter(
                customer_type=
                customer_type
            )
        )


    # =====================================================
    # PAYMENT STATUS
    # =====================================================

    valid_payment_types = {
        value
        for value, label
        in Distributor.PAYMENT_TYPES
    }

    if payment_status in valid_payment_types:

        distributors = (
            distributors.filter(
                payment_status=
                payment_status
            )
        )


    # =====================================================
    # TERRITORY
    # =====================================================

    if territory_id.isdigit():

        distributors = (
            distributors.filter(
                territory_id=
                territory_id
            )
        )


    context = {
        "distributors":
            distributors,

        "search":
            search,

        "selected_status":
            status,

        "selected_customer_type":
            customer_type,

        "selected_payment_status":
            payment_status,

        "selected_territory":
            territory_id,

        "customer_type_choices":
            Distributor.CUSTOMER_TYPES,

        "payment_type_choices":
            Distributor.PAYMENT_TYPES,

        "territories":
            Territory.objects
            .all()
            .order_by("name"),
    }


    return render(
        request,
        "orders/distributor_list.html",
        context,
    )


@login_required
def distributor_create(request):

    if request.method == "POST":

        data = request.POST.copy()

        # Generate a Customer ID if the field
        # is somehow submitted empty.
        if not data.get(
            "customer_code"
        ):

            data["customer_code"] = (
                _generate_customer_code()
            )

        form = DistributorForm(
            data
        )

        if form.is_valid():

            distributor = (
                form.save()
            )

            messages.success(
                request,
                (
                    f"Customer "
                    f"{distributor.full_name} "
                    f"({distributor.customer_code}) "
                    f"was created successfully."
                )
            )

            return redirect(
                "orders:distributor_list"
            )

    else:

        form = DistributorForm(
            initial={
                "customer_code":
                    _generate_customer_code(),

                "is_active":
                    True,
            }
        )


    context = {
        "form":
            form,

        "page_title":
            "Add Customer",

        "editing":
            False,
    }


    return render(
        request,
        "orders/distributor_form.html",
        context,
    )

@login_required
def distributor_edit(
    request,
    distributor_id,
):

    distributor = get_object_or_404(
        Distributor.objects.select_related(
            "territory"
        ),
        pk=distributor_id,
    )

    if request.method == "POST":

        form = DistributorForm(
            request.POST,
            instance=distributor,
        )

        if form.is_valid():

            distributor = form.save()

            messages.success(
                request,
                (
                    f"Customer "
                    f"{distributor.full_name} "
                    f"({distributor.customer_code}) "
                    f"was updated successfully."
                )
            )

            return redirect(
                "orders:distributor_list"
            )

    else:

        form = DistributorForm(
            instance=distributor
        )

    context = {
        "form":
            form,

        "distributor":
            distributor,

        "page_title":
            "Edit Customer",

        "editing":
            True,
    }

    return render(
        request,
        "orders/distributor_form.html",
        context,
    )


@login_required
@require_POST
def distributor_toggle_status(request, distributor_id):
    distributor = get_object_or_404(Distributor, pk=distributor_id)

    if not distributor.is_active:
        conflicting = Distributor.objects.filter(
            territory=distributor.territory,
            is_active=True,
        ).exclude(pk=distributor.pk)

        if conflicting.exists():
            messages.error(
                request,
                "This customer cannot be activated because the territory "
                "already has another active distributor."
            )
            return redirect("orders:distributor_list")

    distributor.is_active = not distributor.is_active
    distributor.save(update_fields=["is_active"])

    state = "activated" if distributor.is_active else "deactivated"
    messages.success(request, f"Customer {distributor.full_name} was {state}.")
    return redirect("orders:distributor_list")


# ---------------------------------------------------------
# TERRITORIES
# ---------------------------------------------------------

@login_required
def territory_list(request):

    territories = (
        Territory.objects
        .all()
        .order_by("name")
    )

    search = request.GET.get(
        "q",
        ""
    ).strip()

    status = request.GET.get(
        "status",
        ""
    ).strip()


    # =====================================================
    # SEARCH
    # =====================================================

    if search:

        territories = (
            territories.filter(
                name__icontains=search
            )
        )


    # =====================================================
    # STATUS
    # =====================================================

    if status == "active":

        territories = (
            territories.filter(
                is_active=True
            )
        )

    elif status == "inactive":

        territories = (
            territories.filter(
                is_active=False
            )
        )


    context = {
        "territories":
            territories,

        "search":
            search,

        "selected_status":
            status,
    }


    return render(
        request,
        "orders/territory_list.html",
        context,
    )

@login_required
def territory_create(request):
    if request.method == "POST":
        form = TerritoryForm(request.POST)
        if form.is_valid():
            territory = form.save()
            messages.success(
                request,
                f"Territory {territory.name} was added successfully."
            )
            return redirect("orders:territory_list")
    else:
        form = TerritoryForm(initial={"is_active": True})

    return render(
        request,
        "orders/territory_form.html",
        {"form": form, "page_title": "Add Territory", "editing": False},
    )


@login_required
def territory_edit(request, territory_id):
    territory = get_object_or_404(Territory, pk=territory_id)

    if request.method == "POST":
        form = TerritoryForm(request.POST, instance=territory)
        if form.is_valid():
            territory = form.save()
            messages.success(
                request,
                f"Territory {territory.name} was updated successfully."
            )
            return redirect("orders:territory_list")
    else:
        form = TerritoryForm(instance=territory)

    return render(
        request,
        "orders/territory_form.html",
        {
            "form": form,
            "territory": territory,
            "page_title": "Edit Territory",
            "editing": True,
        },
    )


@login_required
@require_POST
def territory_toggle_status(request, territory_id):
    territory = get_object_or_404(Territory, pk=territory_id)
    territory.is_active = not territory.is_active
    territory.save(update_fields=["is_active"])

    state = "activated" if territory.is_active else "deactivated"
    messages.success(request, f"Territory {territory.name} was {state}.")
    return redirect("orders:territory_list")


# =========================================================
# ORDERS
# =========================================================

@login_required
def order_list(request):

    orders = (
        Order.objects
        .select_related(
            "distributor",
            "distributor__territory"
        )
        .order_by(
            "-order_date",
            "-created_at",
        )
    )

    search = request.GET.get(
        "q",
        ""
    ).strip()

    status = request.GET.get(
        "status",
        ""
    ).strip()

    selected_date = _parse_date(
        request.GET.get("date")
    )

    if search:

        orders = orders.filter(
            Q(
                serial_number__icontains=search
            )
            |
            Q(
                distributor__full_name__icontains=
                search
            )
            |
            Q(
                distributor__customer_code__icontains=
                search
            )
            |
            Q(
                distributor__territory__name__icontains=
                search
            )
        )

    if status:
        orders = orders.filter(
            status=status
        )

    if selected_date:
        orders = orders.filter(
            order_date=selected_date
        )

    context = {
        "orders": orders,
        "search": search,
        "selected_status": status,
        "selected_date": selected_date,
        "status_choices":
            Order.STATUS_CHOICES,
    }

    return render(
        request,
        "orders/order_list.html",
        context,
    )


@login_required
def order_create(request):

    if request.method == "POST":

        data = request.POST.copy()

        # Generate a serial automatically if
        # the clerk leaves it blank.
        if not data.get("serial_number"):

            selected_date = _parse_date(
                data.get("order_date")
            )

            if not selected_date:
                selected_date = (
                    timezone.localdate()
                )

            data["serial_number"] = (
                _generate_order_serial(
                    selected_date
                )
            )

        form = OrderForm(data)

        formset = OrderItemFormSet(
            data,
            instance=form.instance,
        )

        if (
            form.is_valid()
            and formset.is_valid()
        ):

            with transaction.atomic():

                order = form.save()

                formset.instance = order
                formset.save()

            messages.success(
                request,
                (
                    f"Order "
                    f"{order.serial_number} "
                    f"was created successfully."
                )
            )

            return redirect(
                "orders:order_detail",
                order_id=order.id,
            )

    else:

        today = timezone.localdate()

        initial = {
            "order_date": today,
            "serial_number":
                _generate_order_serial(
                    today
                ),
            "status": "DRAFT",
        }

        form = OrderForm(
            initial=initial
        )

        formset = OrderItemFormSet()

    context = {
        "form": form,
        "formset": formset,
        "page_title": "Create Order",
        "editing": False,
    }

    return render(
        request,
        "orders/order_form.html",
        context,
    )


@login_required
def order_detail(request, order_id):

    order = get_object_or_404(
        Order.objects
        .select_related(
            "distributor",
            "distributor__territory"
        )
        .prefetch_related(
            "items__product",
            "dispatches__driver",
            "dispatches__truck",
            "dispatches__items__order_item__product",
        ),
        pk=order_id,
    )

    purchase_request = (
        PurchaseRequest.objects
        .filter(order=order)
        .first()
    )

    context = {
        "order": order,
        "purchase_request":
            purchase_request,
    }

    return render(
        request,
        "orders/order_detail.html",
        context,
    )


@login_required
def order_edit(request, order_id):

    order = get_object_or_404(
        Order,
        pk=order_id,
    )

    if request.method == "POST":

        form = OrderForm(
            request.POST,
            instance=order,
        )

        formset = OrderItemFormSet(
            request.POST,
            instance=order,
        )

        if (
            form.is_valid()
            and formset.is_valid()
        ):

            with transaction.atomic():

                order = form.save()

                formset.instance = order
                formset.save()

            messages.success(
                request,
                (
                    f"Order "
                    f"{order.serial_number} "
                    f"was updated successfully."
                )
            )

            return redirect(
                "orders:order_detail",
                order_id=order.id,
            )

    else:

        form = OrderForm(
            instance=order
        )

        formset = OrderItemFormSet(
            instance=order
        )

    context = {
        "form": form,
        "formset": formset,
        "order": order,
        "page_title": "Edit Order",
        "editing": True,
    }

    return render(
        request,
        "orders/order_form.html",
        context,
    )


@login_required
def order_print(request, order_id):

    order = get_object_or_404(
        Order.objects
        .select_related(
            "distributor",
            "distributor__territory"
        )
        .prefetch_related(
            "items__product"
        ),
        pk=order_id,
    )

    return render(
        request,
        "orders/order_print.html",
        {
            "order": order,
        },
    )


# =========================================================
# DISPATCHES
# =========================================================

@login_required
def dispatch_list(request):

    dispatches = (
        Dispatch.objects
        .select_related(
            "order",
            "order__distributor",
            "order__distributor__territory",
            "driver",
            "truck",
        )
        .prefetch_related(
            "items__order_item__product"
        )
        .order_by(
            "-dispatch_date",
            "-created_at",
        )
    )

    search = request.GET.get(
        "q",
        ""
    ).strip()

    selected_date = _parse_date(
        request.GET.get("date")
    )

    # =====================================================
    # SEARCH
    # =====================================================

    if search:

        dispatches = (
            dispatches.filter(

                Q(
                    dispatch_number__icontains=
                    search
                )

                |

                Q(
                    order__serial_number__icontains=
                    search
                )

                |

                Q(
                    order__distributor__full_name__icontains=
                    search
                )

                |

                Q(
                    order__distributor__customer_code__icontains=
                    search
                )

                |

                Q(
                    order__distributor__territory__name__icontains=
                    search
                )

                |

                Q(
                    driver__full_name__icontains=
                    search
                )

                |

                Q(
                    truck__plate_number__icontains=
                    search
                )

                |

                Q(
                    external_driver_name__icontains=
                    search
                )

                |

                Q(
                    external_truck_plate__icontains=
                    search
                )
            )
        )

    # =====================================================
    # DISPATCH DATE
    # =====================================================

    if selected_date:

        dispatches = (
            dispatches.filter(
                dispatch_date=
                selected_date
            )
        )

    context = {
        "dispatches":
            dispatches,

        "search":
            search,

        "selected_date":
            selected_date,
    }

    return render(
        request,
        "orders/dispatch_list.html",
        context,
    )

@login_required
def dispatch_create(
    request,
    order_id,
):

    order = get_object_or_404(
        Order.objects
        .select_related(
            "distributor",
            "distributor__territory",
        )
        .prefetch_related(
            "items__product"
        ),
        pk=order_id,
    )

    if order.status == "CANCELLED":

        messages.error(
            request,
            "A cancelled order cannot be dispatched."
        )

        return redirect(
            "orders:order_detail",
            order_id=order.id,
        )

    # =====================================================
    # REMAINING ORDER QUANTITIES
    # =====================================================

    order_item_statuses = (
        _get_dispatch_order_item_statuses(
            order
        )
    )

    remaining_items = [
        status

        for status
        in order_item_statuses

        if status["available"] > 0
    ]

    # Do not allow another Dispatch when the
    # entire Order has already been dispatched.
    if not remaining_items:

        messages.warning(
            request,
            (
                "All quantities in this order "
                "have already been dispatched."
            )
        )

        return redirect(
            "orders:order_detail",
            order_id=order.id,
        )

    dispatch = Dispatch(
        order=order
    )

    # =====================================================
    # POST
    # =====================================================

    if request.method == "POST":

        form = DispatchForm(
            request.POST,
            instance=dispatch,
        )

        # The Order comes from the URL and cannot
        # be changed while creating the Dispatch.
        form.fields["order"].queryset = (
            Order.objects.filter(
                pk=order.pk
            )
        )

        form.fields["order"].disabled = True

        formset = DispatchItemFormSet(
            request.POST,
            instance=dispatch,
            form_kwargs={
                "order": order,
            },
        )

        _attach_dispatch_form_metadata(
            formset,
            order_item_statuses,
        )

        forms_valid = (
            form.is_valid()
            and formset.is_valid()
        )

        if forms_valid:

            quantities_valid = (
                _validate_dispatch_quantities(
                    formset,
                    order,
                )
            )

            if quantities_valid:

                with transaction.atomic():

                    dispatch = form.save()

                    formset.instance = (
                        dispatch
                    )

                    formset.save()

                    PurchaseRequest.objects.get_or_create(
                        order=order
                    )

                messages.success(
                    request,
                    (
                        f"Dispatch "
                        f"{dispatch.dispatch_number} "
                        f"was created successfully."
                    )
                )

                return redirect(
                    "orders:dispatch_detail",
                    dispatch_id=dispatch.id,
                )

    # =====================================================
    # GET
    # =====================================================

    else:

        dispatch.dispatch_number = (
            _generate_dispatch_number(
                order
            )
        )

        form = DispatchForm(
            instance=dispatch
        )

        form.fields["order"].queryset = (
            Order.objects.filter(
                pk=order.pk
            )
        )

        form.fields["order"].disabled = True

        # -------------------------------------------------
        # Automatically load every remaining product.
        #
        # The default quantity is the full quantity that
        # still remains to be dispatched.
        #
        # The clerk only needs to reduce quantities for
        # partial dispatches.
        # -------------------------------------------------

        initial_items = [
            {
                "order_item":
                    status[
                        "order_item"
                    ].pk,

                "quantity":
                    status[
                        "available"
                    ],
            }

            for status
            in remaining_items
        ]

        formset = DispatchItemFormSet(
            instance=dispatch,
            initial=initial_items,
            form_kwargs={
                "order": order,
            },
        )

        # DispatchItemFormSet currently declares extra=5.
        #
        # For this page we only want exactly the number
        # of remaining products, not five empty rows.
        #
        # Django builds the forms lazily, so setting
        # extra here controls the GET form count.
        formset.extra = len(
            initial_items
        )

        _attach_dispatch_form_metadata(
            formset,
            order_item_statuses,
        )

    context = {
        "order": order,
        "form": form,
        "formset": formset,
        "order_item_statuses":
            order_item_statuses,
        "page_title":
            "Create Dispatch",
        "editing": False,
    }

    return render(
        request,
        "orders/dispatch_form.html",
        context,
    )


@login_required
def dispatch_detail(
    request,
    dispatch_id,
):

    dispatch = get_object_or_404(
        Dispatch.objects
        .select_related(
            "order",
            "order__distributor",
            "order__distributor__territory",
            "driver",
            "truck",
        )
        .prefetch_related(
            "items__order_item__product"
        ),
        pk=dispatch_id,
    )

    delivery_note = (
        DeliveryNote.objects
        .filter(dispatch=dispatch)
        .first()
    )

    context = {
        "dispatch": dispatch,
        "delivery_note":
            delivery_note,
    }

    return render(
        request,
        "orders/dispatch_detail.html",
        context,
    )

@login_required
def dispatch_edit(
    request,
    dispatch_id,
):

    dispatch = get_object_or_404(
        Dispatch.objects
        .select_related(
            "order",
            "order__distributor",
            "order__distributor__territory",
        ),
        pk=dispatch_id,
    )

    order = dispatch.order

    order_item_statuses = (
        _get_dispatch_order_item_statuses(
            order,
            dispatch=dispatch,
        )
    )

    # =====================================================
    # POST
    # =====================================================

    if request.method == "POST":

        form = DispatchForm(
            request.POST,
            instance=dispatch,
        )

        form.fields["order"].queryset = (
            Order.objects.filter(
                pk=order.pk
            )
        )

        form.fields["order"].disabled = True

        formset = DispatchItemFormSet(
            request.POST,
            instance=dispatch,
            form_kwargs={
                "order": order,
            },
        )

        _attach_dispatch_form_metadata(
            formset,
            order_item_statuses,
        )

        forms_valid = (
            form.is_valid()
            and formset.is_valid()
        )

        if forms_valid:

            quantities_valid = (
                _validate_dispatch_quantities(
                    formset,
                    order,
                    dispatch=dispatch,
                )
            )

            if quantities_valid:

                with transaction.atomic():

                    dispatch = form.save()

                    formset.instance = (
                        dispatch
                    )

                    formset.save()

                    PurchaseRequest.objects.get_or_create(
                        order=order
                    )

                messages.success(
                    request,
                    (
                        f"Dispatch "
                        f"{dispatch.dispatch_number} "
                        f"was updated successfully."
                    )
                )

                return redirect(
                    "orders:dispatch_detail",
                    dispatch_id=dispatch.id,
                )

    # =====================================================
    # GET
    # =====================================================

    else:

        form = DispatchForm(
            instance=dispatch
        )

        form.fields["order"].queryset = (
            Order.objects.filter(
                pk=order.pk
            )
        )

        form.fields["order"].disabled = True

        formset = DispatchItemFormSet(
            instance=dispatch,
            form_kwargs={
                "order": order,
            },
        )

        # Existing DispatchItems are displayed.
        # Five additional empty rows are unnecessary.
        formset.extra = 0

        _attach_dispatch_form_metadata(
            formset,
            order_item_statuses,
        )

    context = {
        "dispatch": dispatch,
        "order": order,
        "form": form,
        "formset": formset,
        "order_item_statuses":
            order_item_statuses,
        "page_title":
            "Edit Dispatch",
        "editing": True,
    }

    return render(
        request,
        "orders/dispatch_form.html",
        context,
    )

# =========================================================
# DISPATCH PRINT MATRIX
# =========================================================

def _build_dispatch_print_matrix(dispatch):
    """
    Build the fixed ABI dispatch matrix.

    The printed dispatch always keeps the standard ABI
    product columns and packing rows, even when a product
    has no quantity in the current dispatch.
    """

    dispatch_items = list(
        dispatch.items.all()
    )


    # =====================================================
    # STANDARD ABI PRODUCT COLUMNS
    # =====================================================

    standard_product_columns = [
        {
            "key": "cristal water",
            "label": "Cristal Water",
        },
        {
            "key": "pepsi",
            "label": "PEPSI",
        },
        {
            "key": "dew",
            "label": "Dew",
        },
        {
            "key": "mirinda",
            "label": "Mirinda",
        },
        {
            "key": "7up",
            "label": "7UP",
        },
        {
            "key": "b. blast",
            "label": "B. Blast",
        },
        {
            "key": "g. rush",
            "label": "G. Rush",
        },
        {
            "key": "cristal ice",
            "label": "Cristal ICE",
        },
    ]


    # =====================================================
    # STANDARD ABI PACKING ROWS
    # =====================================================

    standard_packing_rows = [
        {
            "key": (
                "PET",
                12,
                "500ml",
            ),
            "label":
                "12 Bottles 500 ml",
        },

        {
            "key": (
                "PET",
                6,
                "1500ml",
            ),
            "label":
                "6 Bottles 1500 ml",
        },

        {
            "key": (
                "CAN",
                24,
                "300ml",
            ),
            "label":
                "24 cans crate, 300 ml",
        },

        {
            "key": (
                "CAN",
                24,
                "250ml",
            ),
            "label":
                "24 cans crate, 250 ml",
        },

        {
            "key": (
                "PET",
                20,
                "300ml",
            ),
            "label":
                "20 Bottles 300 ml",
        },
    ]


    # =====================================================
    # NORMALIZATION
    # =====================================================

    def normalize_text(value):

        return (
            str(value)
            .strip()
            .casefold()
        )


    def normalize_size(value):

        return (
            str(value)
            .strip()
            .lower()
            .replace(" ", "")
        )


    def product_key(product):

        return normalize_text(
            product.name
        )


    def packing_key(product):

        return (
            product.package_type,
            product.units_per_case,
            normalize_size(
                product.size
            ),
        )


    # =====================================================
    # PRODUCT NAME ALIASES
    #
    # These make the print layout tolerant of small
    # differences in Product names inside Management.
    # =====================================================

    product_aliases = {
        "cristal water":
            "cristal water",

        "crystal water":
            "cristal water",

        "pepsi":
            "pepsi",

        "dew":
            "dew",

        "mountain dew":
            "dew",

        "mirinda":
            "mirinda",

        "7up":
            "7up",

        "7 up":
            "7up",

        "b. blast":
            "b. blast",

        "b blast":
            "b. blast",

        "b.blast":
            "b. blast",

        "blast":
            "b. blast",

        "g. rush":
            "g. rush",

        "g rush":
            "g. rush",

        "g.rush":
            "g. rush",

        "rush":
            "g. rush",

        "cristal ice":
            "cristal ice",

        "crystal ice":
            "cristal ice",
    }


    # =====================================================
    # QUANTITIES
    # =====================================================

    quantities = defaultdict(int)


    for item in dispatch_items:

        product = (
            item.order_item.product
        )


        raw_product_key = (
            product_key(product)
        )


        normalized_product_key = (
            product_aliases.get(
                raw_product_key,
                raw_product_key,
            )
        )


        normalized_packing_key = (
            packing_key(product)
        )


        quantities[
            (
                normalized_packing_key,
                normalized_product_key,
            )
        ] += item.quantity


    # =====================================================
    # BUILD MATRIX
    # =====================================================

    packing_rows = []

    total_cases = 0


    for packing in standard_packing_rows:

        row_quantities = []

        row_total = 0


        for column in standard_product_columns:

            quantity = quantities.get(
                (
                    packing["key"],
                    column["key"],
                ),
                0,
            )


            row_quantities.append(
                quantity
            )


            row_total += quantity


        packing_rows.append(
            {
                "packing":
                    packing["label"],

                "quantities":
                    row_quantities,

                "total":
                    row_total,
            }
        )


        total_cases += row_total


    return {
        "product_columns":
            standard_product_columns,

        "packing_rows":
            packing_rows,

        "total_cases":
            total_cases,
    }

@login_required
def dispatch_print(
    request,
    dispatch_id,
):

    dispatch = get_object_or_404(
        Dispatch.objects
        .select_related(
            "order",
            "order__distributor",
            "order__distributor__territory",
            "driver",
            "truck",
        )
        .prefetch_related(
            "items__order_item__product"
        ),
        pk=dispatch_id,
    )


    matrix = (
        _build_dispatch_print_matrix(
            dispatch
        )
    )


    context = {
        "dispatch":
            dispatch,

        "product_columns":
            matrix[
                "product_columns"
            ],

        "packing_rows":
            matrix[
                "packing_rows"
            ],

        "total_cases":
            matrix[
                "total_cases"
            ],

        "copy_names": [
            "SALES COPY",
            "WAREHOUSE COPY",
            "SECURITY COPY",
            "FINANCE COPY",
        ],
    }


    return render(
        request,
        "orders/dispatch_print.html",
        context,
    )

# =========================================================
# PURCHASE REQUEST PRINT HELPERS
# =========================================================


def _round_purchase_amount(value):
    """
    Round financial amounts to whole AFN for printing.
    """

    return Decimal(value).quantize(
        Decimal("1"),
        rounding=ROUND_HALF_UP,
    )


def _format_purchase_amount(value):
    """
    Format a rounded monetary amount without decimals.
    """

    return f"{_round_purchase_amount(value):,.0f}"


def _purchase_number_to_words(value):
    """
    Convert a whole-number AFN amount to English words.
    """

    number = int(value)

    if number == 0:
        return "Zero"

    ones = [
        "",
        "One",
        "Two",
        "Three",
        "Four",
        "Five",
        "Six",
        "Seven",
        "Eight",
        "Nine",
        "Ten",
        "Eleven",
        "Twelve",
        "Thirteen",
        "Fourteen",
        "Fifteen",
        "Sixteen",
        "Seventeen",
        "Eighteen",
        "Nineteen",
    ]

    tens = [
        "",
        "",
        "Twenty",
        "Thirty",
        "Forty",
        "Fifty",
        "Sixty",
        "Seventy",
        "Eighty",
        "Ninety",
    ]


    def under_thousand(amount):

        words = []

        if amount >= 100:

            words.append(
                ones[amount // 100]
            )

            words.append(
                "Hundred"
            )

            amount %= 100


        if amount >= 20:

            words.append(
                tens[amount // 10]
            )

            amount %= 10


        if amount:

            words.append(
                ones[amount]
            )


        return " ".join(words)


    groups = [
        (1_000_000_000, "Billion"),
        (1_000_000, "Million"),
        (1_000, "Thousand"),
    ]


    words = []


    for group_value, group_name in groups:

        if number >= group_value:

            group_amount = (
                number // group_value
            )

            words.append(
                under_thousand(
                    group_amount
                )
            )

            words.append(
                group_name
            )

            number %= group_value


    if number:

        words.append(
            under_thousand(
                number
            )
        )


    return " ".join(words)



def _normalize_purchase_brand(value):

    text = (
        str(value)
        .strip()
        .casefold()
    )


    aliases = {
        "cristal water":
            "cristal water",

        "crystal water":
            "cristal water",

        "pepsi":
            "pepsi",

        "dew":
            "mountain dew",

        "mountain dew":
            "mountain dew",

        "mirinda":
            "mirinda",

        "7up":
            "7up",

        "7 up":
            "7up",

        "b. blast":
            "b. blast",

        "b blast":
            "b. blast",

        "b.blast":
            "b. blast",

        "g. rush":
            "g. rush",

        "g rush":
            "g. rush",

        "g.rush":
            "g. rush",
    }


    return aliases.get(
        text,
        text,
    )



def _normalize_purchase_size(value):

    text = (
        str(value)
        .strip()
        .casefold()
        .replace(" ", "")
    )


    aliases = {
        "0.5litre": "500ml",
        "0.5liter": "500ml",
        "0.5l": "500ml",
        "500ml": "500ml",

        "1.5litre": "1500ml",
        "1.5liter": "1500ml",
        "1.5l": "1500ml",
        "1500ml": "1500ml",

        "250ml": "250ml",
        "300ml": "300ml",
    }


    return aliases.get(
        text,
        text,
    )



def _build_purchase_request_rows(order):
    """
    Build the fixed product table used by ABI's
    existing Purchase Request.

    All standard products are printed even when their
    order quantity is zero.
    """

    standard_rows = [

        {
            "brand_key": "cristal water",
            "brand": "Cristal Water",
            "size_key": "500ml",
            "pack": "0.5 Litre",
            "package_type": "PET",
            "shade": False,
            "emphasis": False,
        },

        {
            "brand_key": "cristal water",
            "brand": "Cristal Water",
            "size_key": "1500ml",
            "pack": "1.5 Litre",
            "package_type": "PET",
            "shade": True,
            "emphasis": True,
        },

        {
            "brand_key": "pepsi",
            "brand": "Pepsi",
            "size_key": "250ml",
            "pack": "250 ml",
            "package_type": "CAN",
            "shade": False,
            "emphasis": False,
        },

        {
            "brand_key": "pepsi",
            "brand": "Pepsi",
            "size_key": "300ml",
            "pack": "300 ml",
            "package_type": "CAN",
            "shade": True,
            "emphasis": False,
        },

        {
            "brand_key": "pepsi",
            "brand": "Pepsi",
            "size_key": "1500ml",
            "pack": "1.5 Litre",
            "package_type": "PET",
            "shade": True,
            "emphasis": True,
        },

        {
            "brand_key": "mountain dew",
            "brand": "Mountain Dew",
            "size_key": "250ml",
            "pack": "250 ml",
            "package_type": "CAN",
            "shade": False,
            "emphasis": False,
        },

        {
            "brand_key": "mountain dew",
            "brand": "Mountain Dew",
            "size_key": "300ml",
            "pack": "300 ml",
            "package_type": "CAN",
            "shade": True,
            "emphasis": False,
        },

        {
            "brand_key": "mountain dew",
            "brand": "Mountain Dew",
            "size_key": "1500ml",
            "pack": "1.5 Litre",
            "package_type": "PET",
            "shade": True,
            "emphasis": True,
        },

        {
            "brand_key": "mirinda",
            "brand": "Mirinda",
            "size_key": "250ml",
            "pack": "250 ml",
            "package_type": "CAN",
            "shade": False,
            "emphasis": False,
        },

        {
            "brand_key": "mirinda",
            "brand": "Mirinda",
            "size_key": "300ml",
            "pack": "300 ml",
            "package_type": "CAN",
            "shade": True,
            "emphasis": False,
        },

        {
            "brand_key": "mirinda",
            "brand": "Mirinda",
            "size_key": "1500ml",
            "pack": "1.5 Litre",
            "package_type": "PET",
            "shade": True,
            "emphasis": True,
        },

        {
            "brand_key": "7up",
            "brand": "7UP",
            "size_key": "250ml",
            "pack": "250 ml",
            "package_type": "CAN",
            "shade": False,
            "emphasis": False,
        },

        {
            "brand_key": "7up",
            "brand": "7UP",
            "size_key": "300ml",
            "pack": "300 ml",
            "package_type": "CAN",
            "shade": True,
            "emphasis": False,
        },

        {
            "brand_key": "7up",
            "brand": "7UP",
            "size_key": "1500ml",
            "pack": "1.5 Litre",
            "package_type": "PET",
            "shade": True,
            "emphasis": True,
        },

        {
            "brand_key": "b. blast",
            "brand": "B. Blast",
            "size_key": "250ml",
            "pack": "250 ml",
            "package_type": "CAN",
            "shade": False,
            "emphasis": False,
        },

        {
            "brand_key": "g. rush",
            "brand": "G. Rush",
            "size_key": "250ml",
            "pack": "250 ml",
            "package_type": "CAN",
            "shade": True,
            "emphasis": True,
        },

        {
            "brand_key": "b. blast",
            "brand": "B. Blast",
            "size_key": "300ml",
            "pack": "300 ml PET",
            "package_type": "PET",
            "shade": False,
            "emphasis": False,
        },

        {
            "brand_key": "g. rush",
            "brand": "G. Rush",
            "size_key": "300ml",
            "pack": "300 ml PET",
            "package_type": "PET",
            "shade": True,
            "emphasis": True,
        },
    ]


    # =====================================================
    # PRODUCT MASTER LOOKUP
    # =====================================================

    product_lookup = {}


    products = (
        Product.objects
        .all()
        .order_by(
            "-is_active",
            "name",
            "size",
        )
    )


    for product in products:

        key = (
            _normalize_purchase_brand(
                product.name
            ),
            _normalize_purchase_size(
                product.size
            ),
            product.package_type,
        )


        product_lookup.setdefault(
            key,
            product,
        )


    # =====================================================
    # ORDER ITEM LOOKUP
    # =====================================================

    order_item_lookup = {}


    for item in order.items.all():

        key = (
            _normalize_purchase_brand(
                item.product.name
            ),
            _normalize_purchase_size(
                item.product.size
            ),
            item.product.package_type,
        )


        order_item_lookup[
            key
        ] = item


    # =====================================================
    # BUILD PRINT ROWS
    # =====================================================

    rows = []


    for definition in standard_rows:

        key = (
            definition[
                "brand_key"
            ],
            definition[
                "size_key"
            ],
            definition[
                "package_type"
            ],
        )


        order_item = (
            order_item_lookup.get(
                key
            )
        )


        product = (
            product_lookup.get(
                key
            )
        )


        if order_item:

            quantity = (
                order_item.quantity
            )

            price = (
                order_item.unit_price
            )

        else:

            quantity = 0


            if product:

                if (
                    order.distributor.customer_type
                    == "WHOLESALER"
                ):

                    price = (
                        product.wholesale_price
                    )

                else:

                    price = (
                        product.retail_price
                    )

            else:

                price = None


        if price is not None:

            price_display = (
                _format_purchase_amount(
                    price
                )
            )

        else:

            price_display = "-"


        if (
            quantity
            and price is not None
        ):

            line_amount = (
                Decimal(quantity)
                * price
            )

            amount_display = (
                _format_purchase_amount(
                    line_amount
                )
            )

        else:

            amount_display = "-"


        rows.append(
            {
                "brand":
                    definition[
                        "brand"
                    ],

                "pack":
                    definition[
                        "pack"
                    ],

                "quantity":
                    quantity,

                "price_display":
                    price_display,

                "amount_display":
                    amount_display,

                "shade":
                    definition[
                        "shade"
                    ],

                "emphasis":
                    definition[
                        "emphasis"
                    ],
            }
        )


    return rows



def _join_purchase_values(values):
    """
    Join unique dispatch values without repeating them.
    """

    result = []

    seen = set()


    for value in values:

        value = (
            str(value).strip()
            if value
            else "-"
        )


        if value not in seen:

            seen.add(value)

            result.append(value)


    return " / ".join(result) if result else "-"


# =========================================================
# PURCHASE REQUEST
# =========================================================

@login_required
def purchase_request_detail(
    request,
    order_id,
):

    order = get_object_or_404(
        Order.objects
        .select_related(
            "distributor",
            "distributor__territory"
        )
        .prefetch_related(
            "items__product",
            "dispatches__items__order_item__product",
            "dispatches__driver",
            "dispatches__truck",
        ),
        pk=order_id,
    )

    purchase_request, created = (
        PurchaseRequest.objects
        .get_or_create(
            order=order
        )
    )

    context = {
        "purchase_request":
            purchase_request,
        "order": order,
    }

    return render(
        request,
        "orders/purchase_request_detail.html",
        context,
    )


@login_required
def purchase_request_print(
    request,
    order_id,
):

    order = get_object_or_404(
        Order.objects
        .select_related(
            "distributor",
            "distributor__territory"
        )
        .prefetch_related(
            "items__product",
            "dispatches__items__order_item__product",
            "dispatches__driver",
            "dispatches__truck",
        ),
        pk=order_id,
    )


    purchase_request, created = (
        PurchaseRequest.objects
        .get_or_create(
            order=order
        )
    )


    dispatches = list(
        order.dispatches.all()
    )


    dispatch_numbers = (
        _join_purchase_values(
            dispatch.dispatch_number
            for dispatch in dispatches
        )
    )


    vehicle_numbers = (
        _join_purchase_values(
            (
                dispatch.truck.plate_number
                if dispatch.truck

                else
                dispatch.external_truck_plate
            )
            for dispatch in dispatches
        )
    )


    driver_names = (
        _join_purchase_values(
            (
                dispatch.driver.full_name
                if dispatch.driver

                else
                dispatch.external_driver_name
            )
            for dispatch in dispatches
        )
    )


    subtotal = (
        _round_purchase_amount(
            purchase_request.subtotal
        )
    )


    freight = (
        _round_purchase_amount(
            purchase_request.total_freight
        )
    )


    payable = (
        _round_purchase_amount(
            purchase_request.payable_amount
        )
    )


    context = {
        "purchase_request":
            purchase_request,

        "order":
            order,

        "purchase_rows":
            _build_purchase_request_rows(
                order
            ),

        "dispatch_numbers":
            dispatch_numbers,

        "vehicle_numbers":
            vehicle_numbers,

        "driver_names":
            driver_names,

        "total_cases_display":
            f"{order.total_cases:,}",

        "total_ton_display":
            (
                f"{purchase_request.total_weight_tons:,.2f}"
            ),

        "subtotal_display":
            f"{subtotal:,.0f}",

        "freight_display":
            f"{freight:,.0f}",

        "payable_display":
            f"{payable:,.0f}",

        # The existing ABI document writes
        # the SUBTOTAL in words.
        "amount_words":
            (
                _purchase_number_to_words(
                    subtotal
                )
                + " AFN Only"
            ),
    }


    return render(
        request,
        "orders/purchase_request_print.html",
        context,
    )


# =========================================================
# DELIVERY NOTE
# =========================================================


def _get_delivery_note_dispatch(dispatch_id):
    """
    Load a Dispatch with everything required by the
    Delivery Note pages and printable document.
    """

    return get_object_or_404(
        Dispatch.objects
        .select_related(
            "order",
            "order__distributor",
            "order__distributor__territory",
            "driver",
            "truck",
        )
        .prefetch_related(
            "items__order_item__product"
        ),
        pk=dispatch_id,
    )



def _normalize_delivery_note_brand(value):
    """
    Normalize product names so they can be mapped to
    ABI's fixed Delivery Note product list.
    """

    text = (
        str(value)
        .strip()
        .casefold()
        .replace(".", "")
    )

    text = " ".join(
        text.split()
    )


    if (
        "cristal" in text
        or
        "crystal" in text
    ):
        return "cristal water"


    if "pepsi" in text:
        return "pepsi"


    if (
        text == "dew"
        or
        "mountain dew" in text
    ):
        return "dew"


    if "mirinda" in text:
        return "mirinda"


    if (
        "7up" in text
        or
        "7 up" in text
    ):
        return "7up"


    if (
        "b blast" in text
        or
        text == "blast"
    ):
        return "b blast"


    if (
        "g rush" in text
        or
        text == "rush"
    ):
        return "g rush"


    return text



def _normalize_delivery_note_size(value):
    """
    Normalize different product size formats.
    """

    text = (
        str(value)
        .strip()
        .casefold()
        .replace(" ", "")
    )


    aliases = {
        "0.5l": "500ml",
        "0.5litre": "500ml",
        "0.5liter": "500ml",
        "500ml": "500ml",

        "1.5l": "1500ml",
        "1.5litre": "1500ml",
        "1.5liter": "1500ml",
        "1500ml": "1500ml",

        "250ml": "250ml",
        "300ml": "300ml",
    }


    return aliases.get(
        text,
        text,
    )



def _delivery_note_product_key(product):
    """
    Create a normalized lookup key for Delivery Note
    product quantities.
    """

    return (
        _normalize_delivery_note_brand(
            product.name
        ),
        _normalize_delivery_note_size(
            product.size
        ),
        product.units_per_case,
    )



def _build_delivery_note_rows(dispatch):
    """
    Build the fixed product list used by ABI's
    current Delivery Note.

    All standard rows appear even when their quantity
    for this Dispatch is zero.
    """

    standard_rows = [

        {
            "brand_key": "cristal water",
            "item_name": "Cristal Bottled Mineral Water",
            "size_key": "500ml",
            "units_per_case": 12,
            "description": "500 ml x 12",
        },

        {
            "brand_key": "cristal water",
            "item_name": "Cristal Bottled Mineral Water",
            "size_key": "1500ml",
            "units_per_case": 6,
            "description": "1500 ml x 6",
        },

        {
            "brand_key": "pepsi",
            "item_name": "Pepsi",
            "size_key": "250ml",
            "units_per_case": 24,
            "description": "250 ml x 24",
        },

        {
            "brand_key": "pepsi",
            "item_name": "Pepsi",
            "size_key": "300ml",
            "units_per_case": 24,
            "description": "300 ml x 24",
        },

        {
            "brand_key": "pepsi",
            "item_name": "Pepsi",
            "size_key": "1500ml",
            "units_per_case": 6,
            "description": "1500 ml x 6",
        },

        {
            "brand_key": "dew",
            "item_name": "Dew",
            "size_key": "250ml",
            "units_per_case": 24,
            "description": "250 ml x 24",
        },

        {
            "brand_key": "dew",
            "item_name": "Dew",
            "size_key": "300ml",
            "units_per_case": 24,
            "description": "300 ml x 24",
        },

        {
            "brand_key": "dew",
            "item_name": "Dew",
            "size_key": "1500ml",
            "units_per_case": 6,
            "description": "1500 ml x 6",
        },

        {
            "brand_key": "mirinda",
            "item_name": "Mirinda",
            "size_key": "250ml",
            "units_per_case": 24,
            "description": "250 ml x 24",
        },

        {
            "brand_key": "mirinda",
            "item_name": "Mirinda",
            "size_key": "300ml",
            "units_per_case": 24,
            "description": "300 ml x 24",
        },

        {
            "brand_key": "mirinda",
            "item_name": "Mirinda",
            "size_key": "1500ml",
            "units_per_case": 6,
            "description": "1500 ml x 6",
        },

        {
            "brand_key": "7up",
            "item_name": "7Up",
            "size_key": "250ml",
            "units_per_case": 24,
            "description": "250 ml x 24",
        },

        {
            "brand_key": "7up",
            "item_name": "7Up",
            "size_key": "300ml",
            "units_per_case": 24,
            "description": "300 ml x 24",
        },

        {
            "brand_key": "7up",
            "item_name": "7Up",
            "size_key": "1500ml",
            "units_per_case": 6,
            "description": "1500 ml x 6",
        },

        {
            "brand_key": "b blast",
            "item_name": "B. Blast",
            "size_key": "250ml",
            "units_per_case": 24,
            "description": "250 ml x 24",
        },

        {
            "brand_key": "g rush",
            "item_name": "G. Rush",
            "size_key": "250ml",
            "units_per_case": 24,
            "description": "250 ml x 24",
        },
    ]


    quantities = defaultdict(int)


    for dispatch_item in dispatch.items.all():

        product = (
            dispatch_item
            .order_item
            .product
        )


        key = (
            _delivery_note_product_key(
                product
            )
        )


        quantities[key] += (
            dispatch_item.quantity
        )


    rows = []


    for definition in standard_rows:

        key = (
            definition["brand_key"],
            definition["size_key"],
            definition["units_per_case"],
        )


        rows.append(
            {
                "item_name":
                    definition["item_name"],

                "description":
                    definition["description"],

                "quantity":
                    quantities.get(
                        key,
                        0,
                    ),
            }
        )


    return rows



# =========================================================
# CREATE DELIVERY NOTE
# =========================================================

@login_required
def delivery_note_create(
    request,
    dispatch_id,
):

    dispatch = (
        _get_delivery_note_dispatch(
            dispatch_id
        )
    )


    # Creating a Delivery Note must be a deliberate action.
    # Simply opening a page must not create one.
    if request.method != "POST":

        return redirect(
            "orders:dispatch_detail",
            dispatch_id=dispatch.id,
        )


    delivery_note, created = (
        DeliveryNote.objects
        .get_or_create(
            dispatch=dispatch
        )
    )


    if created:

        messages.success(
            request,
            (
                "Delivery Note created successfully. "
                "Two copies will be prepared when printed."
            )
        )

    else:

        messages.info(
            request,
            (
                "A Delivery Note already exists "
                "for this dispatch."
            )
        )


    return redirect(
        "orders:delivery_note_detail",
        dispatch_id=dispatch.id,
    )



# =========================================================
# DELIVERY NOTE DETAIL
# =========================================================

@login_required
def delivery_note_detail(
    request,
    dispatch_id,
):

    dispatch = (
        _get_delivery_note_dispatch(
            dispatch_id
        )
    )


    delivery_note = get_object_or_404(
        DeliveryNote,
        dispatch=dispatch,
    )


    context = {
        "dispatch":
            dispatch,

        "delivery_note":
            delivery_note,
    }


    return render(
        request,
        "orders/delivery_note_detail.html",
        context,
    )



# =========================================================
# EDIT / RECORD RETURNED COPY
# =========================================================

@login_required
def delivery_note_edit(
    request,
    dispatch_id,
):

    dispatch = (
        _get_delivery_note_dispatch(
            dispatch_id
        )
    )


    delivery_note = get_object_or_404(
        DeliveryNote,
        dispatch=dispatch,
    )


    if request.method == "POST":

        form = DeliveryNoteForm(
            request.POST,
            instance=delivery_note,
        )


        if form.is_valid():

            form.save()


            messages.success(
                request,
                (
                    "Returned Delivery Note information "
                    "was saved successfully."
                )
            )


            return redirect(
                "orders:delivery_note_detail",
                dispatch_id=dispatch.id,
            )


    else:

        form = DeliveryNoteForm(
            instance=delivery_note
        )


    context = {
        "dispatch":
            dispatch,

        "delivery_note":
            delivery_note,

        "form":
            form,
    }


    return render(
        request,
        "orders/delivery_note_form.html",
        context,
    )



# =========================================================
# PRINT DELIVERY NOTE
# =========================================================

@login_required
def delivery_note_print(
    request,
    dispatch_id,
):

    dispatch = (
        _get_delivery_note_dispatch(
            dispatch_id
        )
    )


    delivery_note = get_object_or_404(
        DeliveryNote,
        dispatch=dispatch,
    )


    delivery_rows = (
        _build_delivery_note_rows(
            dispatch
        )
    )


    # =====================================================
    # DELIVERY REPRESENTATIVE
    # =====================================================

    if dispatch.driver:

        delivery_representative = (
            dispatch.driver.full_name
        )


    elif dispatch.external_driver_name:

        delivery_representative = (
            dispatch.external_driver_name
        )


    else:

        delivery_representative = ""



    # =====================================================
    # SHIPMENT FACILITY / VEHICLE
    # =====================================================

    if dispatch.truck:

        shipment_facility = (
            dispatch.truck.plate_number
        )


    elif dispatch.external_truck_plate:

        shipment_facility = (
            dispatch.external_truck_plate
        )


    else:

        shipment_facility = ""



    context = {
        "dispatch":
            dispatch,

        "delivery_note":
            delivery_note,

        "delivery_rows":
            delivery_rows,

        "delivery_representative":
            delivery_representative,

        "shipment_facility":
            shipment_facility,

        # Delivery Notes are always printed in two copies.
        "delivery_note_copies":
            range(2),
    }


    return render(
        request,
        "orders/delivery_note_print.html",
        context,
    )

# =========================================================
# DAILY DISPATCH SUMMARY
# =========================================================


def _build_dispatch_summary(selected_date):
    """
    Build the Daily Dispatch Summary for one dispatch date.

    This helper is shared by:
    - the Dispatch Summary webpage
    - the Excel export

    Keeping the calculations in one place ensures that the
    webpage and exported spreadsheet always show the same data.
    """

    dispatches = (
        Dispatch.objects
        .filter(
            dispatch_date=selected_date
        )
        .select_related(
            "order",
            "order__distributor",
            "order__distributor__territory",
        )
        .prefetch_related(
            "items__order_item__product"
        )
        .order_by(
            "order__distributor__territory__name",
            "order__distributor__full_name",
        )
    )

    rows = {}

    for dispatch in dispatches:

        distributor = (
            dispatch.order.distributor
        )

        territory = (
            distributor.territory
        )

        # One summary row per distributor / territory
        # for the selected dispatch date.
        key = (
            distributor.id,
            territory.id,
        )

        if key not in rows:

            rows[key] = {
                "customer_name":
                    distributor.full_name,

                "customer_code":
                    distributor.customer_code,

                "territory":
                    territory.name,

                "total_trucks": 0,

                "tonnage":
                    Decimal("0"),

                "pet_1500": 0,
                "can_300": 0,
                "can_250": 0,
                "sting_250": 0,
                "sting_300": 0,

                "total_cases": 0,
            }

        row = rows[key]

        # Every Dispatch represents one truck / trip.
        row["total_trucks"] += 1

        row["tonnage"] += (
            dispatch.total_weight_tons
        )

        for dispatch_item in dispatch.items.all():

            product = (
                dispatch_item
                .order_item
                .product
            )

            quantity = (
                dispatch_item.quantity
            )

            row["total_cases"] += (
                quantity
            )

            product_name = (
                product.name
                .lower()
            )

            product_size = (
                product.size
                .lower()
                .replace(" ", "")
            )

            is_sting = (
                "sting" in product_name
            )

            # ---------------------------------------------
            # PET 1500
            # ---------------------------------------------

            if (
                product.package_type == "PET"
                and "1500" in product_size
                and not is_sting
            ):
                row["pet_1500"] += (
                    quantity
                )

            # ---------------------------------------------
            # CAN 300
            # ---------------------------------------------

            if (
                product.package_type == "CAN"
                and "300" in product_size
                and not is_sting
            ):
                row["can_300"] += (
                    quantity
                )

            # ---------------------------------------------
            # CAN 250
            # ---------------------------------------------

            if (
                product.package_type == "CAN"
                and "250" in product_size
                and not is_sting
            ):
                row["can_250"] += (
                    quantity
                )

            # ---------------------------------------------
            # STING 250
            # ---------------------------------------------

            if (
                is_sting
                and "250" in product_size
            ):
                row["sting_250"] += (
                    quantity
                )

            # ---------------------------------------------
            # STING 300
            # ---------------------------------------------

            if (
                is_sting
                and "300" in product_size
            ):
                row["sting_300"] += (
                    quantity
                )

    summary_rows = list(
        rows.values()
    )

    totals = {
        "total_trucks": 0,
        "tonnage": Decimal("0"),
        "pet_1500": 0,
        "can_300": 0,
        "can_250": 0,
        "sting_250": 0,
        "sting_300": 0,
        "total_cases": 0,
    }

    for row in summary_rows:

        totals["total_trucks"] += (
            row["total_trucks"]
        )

        totals["tonnage"] += (
            row["tonnage"]
        )

        totals["pet_1500"] += (
            row["pet_1500"]
        )

        totals["can_300"] += (
            row["can_300"]
        )

        totals["can_250"] += (
            row["can_250"]
        )

        totals["sting_250"] += (
            row["sting_250"]
        )

        totals["sting_300"] += (
            row["sting_300"]
        )

        totals["total_cases"] += (
            row["total_cases"]
        )

    return (
        summary_rows,
        totals,
    )


@login_required
def dispatch_summary(request):

    selected_date = _parse_date(
        request.GET.get("date")
    )

    if not selected_date:

        selected_date = (
            timezone.localdate()
        )

    summary_rows, totals = (
        _build_dispatch_summary(
            selected_date
        )
    )

    context = {
        "selected_date":
            selected_date,

        "rows":
            summary_rows,

        "totals":
            totals,
    }

    return render(
        request,
        "orders/dispatch_summary.html",
        context,
    )


# =========================================================
# DISPATCH SUMMARY - EXCEL EXPORT
# =========================================================


@login_required
def dispatch_summary_excel(request):

    selected_date = _parse_date(
        request.GET.get("date")
    )

    if not selected_date:

        selected_date = (
            timezone.localdate()
        )

    summary_rows, totals = (
        _build_dispatch_summary(
            selected_date
        )
    )


    # =====================================================
    # WORKBOOK
    # =====================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Dispatch Summary"
    )


    # =====================================================
    # PAGE SETTINGS
    # =====================================================

    worksheet.page_setup.orientation = (
        "landscape"
    )

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = (
        True
    )

    worksheet.freeze_panes = (
        "A6"
    )


    # =====================================================
    # STYLES
    # =====================================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F2937",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="DBEAFE",
    )

    white_bold_font = Font(
        color="FFFFFF",
        bold=True,
    )

    title_font = Font(
        bold=True,
        size=16,
    )

    subtitle_font = Font(
        bold=True,
        size=13,
    )

    bold_font = Font(
        bold=True,
    )

    thin_side = Side(
        style="thin",
        color="CBD5E1",
    )

    table_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )


    # =====================================================
    # REPORT HEADER
    # =====================================================

    worksheet.merge_cells(
        "A1:K1"
    )

    worksheet["A1"] = (
        "Afghanistan Beverage Industries Ltd."
    )

    worksheet["A1"].font = (
        title_font
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )


    worksheet.merge_cells(
        "A2:K2"
    )

    worksheet["A2"] = (
        "Daily Dispatch Summary"
    )

    worksheet["A2"].font = (
        subtitle_font
    )

    worksheet["A2"].alignment = Alignment(
        horizontal="center"
    )


    worksheet.merge_cells(
        "A3:K3"
    )

    worksheet["A3"] = (
        selected_date.strftime(
            "%d-%b-%y"
        )
    )

    worksheet["A3"].alignment = Alignment(
        horizontal="center"
    )


    # =====================================================
    # TABLE HEADERS
    # =====================================================

    headers = [
        "S.#",
        "Customer Name",
        "Territory",
        "Total Trucks",
        "Tonnage",
        "PET 1500",
        "Can 300",
        "Can 250",
        "Sting 250",
        "Sting 300",
        "Total Cases",
    ]

    header_row = 5

    for column_number, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column_number,
            value=header,
        )

        cell.fill = (
            header_fill
        )

        cell.font = (
            white_bold_font
        )

        cell.border = (
            table_border
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


    # =====================================================
    # DATA ROWS
    # =====================================================

    current_row = 6

    for index, row in enumerate(
        summary_rows,
        start=1,
    ):

        values = [
            index,

            (
                f"{row['customer_name']} "
                f"({row['customer_code']})"
            ),

            row["territory"],

            row["total_trucks"],

            float(
                row["tonnage"]
            ),

            row["pet_1500"],
            row["can_300"],
            row["can_250"],
            row["sting_250"],
            row["sting_300"],
            row["total_cases"],
        ]

        for column_number, value in enumerate(
            values,
            start=1,
        ):

            cell = worksheet.cell(
                row=current_row,
                column=column_number,
                value=value,
            )

            cell.border = (
                table_border
            )

            if column_number == 2:

                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                    wrap_text=True,
                )

            else:

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        # Tonnage
        worksheet.cell(
            row=current_row,
            column=5,
        ).number_format = (
            "0.00"
        )

        current_row += 1


    # =====================================================
    # TOTAL ROW
    # =====================================================

    worksheet.merge_cells(
        start_row=current_row,
        start_column=1,
        end_row=current_row,
        end_column=3,
    )

    total_label = worksheet.cell(
        row=current_row,
        column=1,
        value="TOTAL",
    )

    total_label.font = (
        bold_font
    )

    total_label.fill = (
        total_fill
    )

    total_label.border = (
        table_border
    )

    total_label.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


    total_values = [
        totals["total_trucks"],
        float(
            totals["tonnage"]
        ),
        totals["pet_1500"],
        totals["can_300"],
        totals["can_250"],
        totals["sting_250"],
        totals["sting_300"],
        totals["total_cases"],
    ]


    for column_number, value in enumerate(
        total_values,
        start=4,
    ):

        cell = worksheet.cell(
            row=current_row,
            column=column_number,
            value=value,
        )

        cell.font = (
            bold_font
        )

        cell.fill = (
            total_fill
        )

        cell.border = (
            table_border
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


    worksheet.cell(
        row=current_row,
        column=5,
    ).number_format = (
        "0.00"
    )


    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    column_widths = {
        1: 7,
        2: 34,
        3: 20,
        4: 13,
        5: 12,
        6: 12,
        7: 11,
        8: 11,
        9: 12,
        10: 12,
        11: 13,
    }


    for column_number, width in (
        column_widths.items()
    ):

        worksheet.column_dimensions[
            get_column_letter(
                column_number
            )
        ].width = width


    # =====================================================
    # EXCEL FILTERS
    # =====================================================

    if summary_rows:

        worksheet.auto_filter.ref = (
            f"A5:K{current_row - 1}"
        )


    # =====================================================
    # CREATE RESPONSE
    # =====================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)


    filename = (
        "Dispatch_Summary_"
        f"{selected_date.strftime('%d-%b-%y')}"
        ".xlsx"
    )


    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )


    return response
